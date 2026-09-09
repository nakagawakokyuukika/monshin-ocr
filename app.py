"""
問診票 OCR アプリ
外部 AI API は使用しない（院内情報管理ポリシー遵守）
"""
import streamlit as st
import json
import io
from datetime import datetime

# --- Tesseract チェック ---
try:
    import pytesseract
    from pytesseract import TesseractNotFoundError
    _TESSERACT_OK = True
except ImportError:
    _TESSERACT_OK = False
    TesseractNotFoundError = Exception

# --- その他ライブラリ ---
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

# ページ設定
st.set_page_config(
    page_title="問診票 OCR",
    page_icon="🏥",
    layout="centered",
    initial_sidebar_state="collapsed",
)

# --- スマホ向け CSS ---
st.markdown("""
<style>
/* 全体フォントサイズ */
html, body, [class*="css"] {
    font-size: 1.1rem !important;
}

/* ボタンを大きく */
div.stButton > button,
div.stDownloadButton > button {
    height: 3.5rem !important;
    font-size: 1.1rem !important;
    width: 100% !important;
    border-radius: 8px !important;
}

/* ファイルアップローダを大きく */
section[data-testid="stFileUploadDropzone"] {
    padding: 2rem 1rem !important;
    font-size: 1.1rem !important;
}

/* 横スクロール防止 */
.main .block-container {
    max-width: 100% !important;
    padding-left: 1rem !important;
    padding-right: 1rem !important;
}

/* テキスト入力 */
input[type="text"] {
    font-size: 1.1rem !important;
    height: 3rem !important;
}
</style>
""", unsafe_allow_html=True)


# ===== ユーティリティ関数 =====

def compress_image(image_bytes: bytes, max_width: int = 1600) -> bytes:
    """画像を最大幅に合わせてリサイズし JPEG に変換する。10MB 超でも処理可能。"""
    img = Image.open(io.BytesIO(image_bytes))
    if img.mode in ("RGBA", "P"):
        img = img.convert("RGB")
    w, h = img.size
    if w > max_width or h > max_width:
        ratio = max_width / max(w, h)
        img = img.resize((int(w * ratio), int(h * ratio)), Image.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=80, optimize=True)
    buf.seek(0)
    return buf.getvalue()


def run_ocr(image_bytes: bytes) -> str:
    """Tesseract で OCR を実行する。"""
    if not _TESSERACT_OK:
        st.error(
            "pytesseract がインポートできません。\n\n"
            "requirements.txt に `pytesseract>=0.3.10` が含まれているか確認してください。"
        )
        st.stop()

    # 圧縮（10MB 超も含め自動対応）
    image_bytes = compress_image(image_bytes)
    img = Image.open(io.BytesIO(image_bytes))

    try:
        custom_config = r"--oem 1 --psm 3"
        text = pytesseract.image_to_string(img, lang="jpn", config=custom_config)
    except TesseractNotFoundError:
        st.error(
            "Tesseract がインストールされていません。\n\n"
            "Streamlit Cloud にデプロイしている場合は、リポジトリに `packages.txt` を追加し、\n"
            "以下の内容を記述してください：\n\n"
            "```\ntesseract-ocr\ntesseract-ocr-jpn\ntesseract-ocr-jpn-vert\n```\n\n"
            "ローカル環境では `sudo apt install tesseract-ocr tesseract-ocr-jpn` を実行してください。"
        )
        st.stop()
    except Exception:
        st.error("OCR 処理中にエラーが発生しました。画像を確認してからやり直してください。")
        st.stop()

    result = text.strip()
    if not result:
        return "（OCRでテキストを取得できませんでした）"
    return result


def build_word_doc(ocr_text: str, patient_no: str, image_bytes: bytes) -> bytes:
    """Word ファイルを生成して bytes で返す。"""
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    now = datetime.now()

    # タイトル
    title = doc.add_heading("問　診　票（OCR）", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 基本情報テーブル
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = "Table Grid"
    cells = tbl.rows[0].cells
    cells[0].text = f"受診日：{now.strftime('%Y年%m月%d日')}"
    cells[1].text = f"患者番号：{patient_no or '─'}"
    cells[2].text = f"処理時刻：{now.strftime('%H:%M')}"
    for cell in cells:
        for para in cell.paragraphs:
            if para.runs:
                para.runs[0].font.size = Pt(10)

    doc.add_paragraph()

    # OCR 読み取り結果
    h2 = doc.add_heading("【 OCR 読み取り結果 】", level=2)
    h2.runs[0].font.size = Pt(12)

    SECTION_KEYWORDS = [
        "どのような症状", "その症状はいつから", "現在治療中", "過去にかかった",
        "現在服用中", "アレルギー", "お酒は", "タバコは", "女性の方", "妊娠", "授乳",
    ]
    for line in ocr_text.splitlines():
        line = line.strip()
        if not line:
            continue
        is_header = any(kw in line for kw in SECTION_KEYWORDS)
        para = doc.add_paragraph()
        run = para.add_run(line)
        run.font.size = Pt(10.5)
        if is_header:
            run.bold = True
            run.font.size = Pt(11)

    doc.add_paragraph()

    # スタッフ確認欄
    h2b = doc.add_heading("【 スタッフ確認欄 】", level=2)
    h2b.runs[0].font.size = Pt(12)
    doc.add_paragraph("確認者：＿＿＿＿＿＿＿　　確認日時：＿＿＿＿＿＿＿")
    doc.add_paragraph()
    doc.add_paragraph("補足・修正：")
    doc.add_paragraph("　" * 80)

    doc.add_paragraph()

    # 原本画像
    h2c = doc.add_heading("【 問診票 原本画像 】", level=2)
    h2c.runs[0].font.size = Pt(12)
    img = Image.open(io.BytesIO(image_bytes))
    if img.mode in ("RGBA", "P"):
        img = img.convert("RGB")
    img_buf = io.BytesIO()
    img.save(img_buf, format="JPEG", quality=85)
    img_buf.seek(0)
    doc.add_picture(img_buf, width=Cm(14))

    out_buf = io.BytesIO()
    doc.save(out_buf)
    out_buf.seek(0)
    return out_buf.getvalue()


@st.cache_resource(show_spinner=False, ttl=300)
def get_drive_service():
    """Google Drive サービスオブジェクトを返す。失敗時は例外を投げる。"""
    from google.oauth2 import service_account
    from googleapiclient.discovery import build

    credentials_info = json.loads(st.secrets["GOOGLE_SERVICE_ACCOUNT"])
    creds = service_account.Credentials.from_service_account_info(
        credentials_info,
        scopes=["https://www.googleapis.com/auth/drive"],
    )
    return build("drive", "v3", credentials=creds, cache_discovery=False)


def save_to_drive(service, doc_bytes: bytes, filename: str) -> str:
    """Drive の日付フォルダにファイルを保存し、WebViewLink を返す。"""
    from googleapiclient.http import MediaIoBaseUpload

    parent_id = st.secrets["DRIVE_FOLDER_ID"]
    today = datetime.now().strftime("%Y-%m-%d")

    # 日付フォルダの検索または作成
    q = (
        f"name='{today}' and '{parent_id}' in parents "
        f"and mimeType='application/vnd.google-apps.folder' and trashed=false"
    )
    results = service.files().list(q=q, fields="files(id)").execute()
    if results["files"]:
        date_folder_id = results["files"][0]["id"]
    else:
        folder_meta = {
            "name": today,
            "mimeType": "application/vnd.google-apps.folder",
            "parents": [parent_id],
        }
        folder = service.files().create(body=folder_meta, fields="id").execute()
        date_folder_id = folder["id"]

    # ファイルアップロード
    file_meta = {"name": filename, "parents": [date_folder_id]}
    media = MediaIoBaseUpload(
        io.BytesIO(doc_bytes),
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        resumable=False,
    )
    uploaded = service.files().create(
        body=file_meta, media_body=media, fields="id,webViewLink"
    ).execute()
    return uploaded.get("webViewLink", "")


# ===== UI =====

st.title("🏥 問診票 OCR")
st.caption("手書き問診票の写真を Word ファイルに変換します")

# 患者番号入力
patient_no = st.text_input(
    "患者番号（任意）",
    placeholder="例：P-12345　※空欄でも可",
    max_chars=30,
)

# ファイルアップロード
uploaded_file = st.file_uploader(
    "📷 問診票の写真を選択",
    type=["jpg", "jpeg", "png"],
    help="JPG または PNG 形式。10MB 超の場合は自動的に圧縮されます。",
)

# 画像プレビュー
if uploaded_file is not None:
    st.image(uploaded_file, caption="アップロードした画像", use_container_width=True)

st.divider()

# OCR 実行ボタン（写真が選択されていない場合は disabled）
ocr_button = st.button(
    "✅ OCR 処理を実行",
    disabled=(uploaded_file is None),
    use_container_width=True,
    type="primary",
)

if ocr_button and uploaded_file is not None:
    image_bytes = uploaded_file.read()

    # ファイル名生成
    now_str = datetime.now().strftime("%Y%m%d_%H%M%S")
    if patient_no.strip():
        # 患者番号に使えない文字を除去
        safe_no = "".join(c for c in patient_no.strip() if c not in r'\/:*?"<>|')
        filename = f"問診票_{safe_no}_{now_str}.docx"
    else:
        filename = f"問診票_{now_str}.docx"

    # OCR 実行
    with st.spinner("OCR 処理中です。しばらくお待ちください..."):
        ocr_text = run_ocr(image_bytes)

    # OCR 結果表示
    st.subheader("OCR 読み取り結果")
    if ocr_text == "（OCRでテキストを取得できませんでした）":
        st.warning("OCRでテキストを取得できませんでした。画像が鮮明か確認してください。")
    else:
        st.text_area("テキスト（参考表示）", value=ocr_text, height=200)

    # Word ファイル生成
    with st.spinner("Word ファイルを生成中..."):
        try:
            doc_bytes = build_word_doc(ocr_text, patient_no.strip(), image_bytes)
        except Exception:
            st.error("Word ファイルの生成に失敗しました。画像を確認してからやり直してください。")
            st.stop()

    # ダウンロードボタン（Drive 結果に関わらず必ず表示）
    st.success("処理が完了しました。")
    st.download_button(
        label="⬇️ Word ファイルをダウンロード",
        data=doc_bytes,
        file_name=filename,
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        use_container_width=True,
        type="primary",
    )

    # Google Drive 保存（任意・失敗してもクラッシュしない）
    drive_enabled = (
        "GOOGLE_SERVICE_ACCOUNT" in st.secrets and "DRIVE_FOLDER_ID" in st.secrets
    )
    if drive_enabled:
        with st.spinner("Google Drive に保存中..."):
            try:
                service = get_drive_service()
                link = save_to_drive(service, doc_bytes, filename)
                if link:
                    st.info(f"Google Drive に保存しました: [ファイルを開く]({link})")
                else:
                    st.info("Google Drive への保存が完了しました。")
            except Exception:
                st.warning(
                    "Google Drive への保存に失敗しました。"
                    "ダウンロードボタンからファイルを取得してください。"
                )
